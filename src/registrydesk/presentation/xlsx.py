"""XLSX exporter for the clean owner view."""

from collections import defaultdict
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from registrydesk.domain.exporting import EXPORT_FIELDS_BY_KEY, validate_export_fields
from registrydesk.domain.contracts import PropertyType
from registrydesk.domain.registry import RegistryDetails, RegistryRow
from registrydesk.presentation.values import owner_display_name


class XlsxExporter:
    """Create a grouped workbook with exact fraction and ownership-area formulas."""

    def export(self, registry: RegistryDetails, field_keys: Iterable[str], path: Path) -> None:
        """Render a grouped workbook while preserving exact ownership formulas."""
        keys = validate_export_fields(field_keys)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self._sheet_title(registry)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False

        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        section_fill = PatternFill("solid", fgColor="EDEDED")

        for column, key in enumerate(keys, start=1):
            cell = sheet.cell(row=1, column=column, value=EXPORT_FIELDS_BY_KEY[key].header_label)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill
            cell.border = border
        sheet.row_dimensions[1].height = 30

        current_row = 2
        sequence = 1
        total_area_data_start: int | None = None
        total_area_data_end: int | None = None

        for section_type, section_title in (
            (PropertyType.APARTMENT.value, "КВАРТИРИ"),
            (PropertyType.PREMISES.value, "НЕЖИТЛОВІ ПРИМІЩЕННЯ"),
        ):
            section_rows = [row for row in registry.rows if row.property_type == section_type]
            if not section_rows:
                continue

            sheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=len(keys),
            )
            section_cell = sheet.cell(row=current_row, column=1, value=section_title)
            section_cell.font = Font(bold=True)
            section_cell.alignment = Alignment(horizontal="left", vertical="center")
            section_cell.fill = section_fill
            for col in range(1, len(keys) + 1):
                sheet.cell(row=current_row, column=col).border = border
                sheet.cell(row=current_row, column=col).fill = section_fill
            current_row += 1

            by_property: dict[int, list[RegistryRow]] = defaultdict(list)
            property_order: list[int] = []
            for row in section_rows:
                if row.property_id not in by_property:
                    property_order.append(row.property_id)
                by_property[row.property_id].append(row)

            for property_id in property_order:
                owner_rows = by_property[property_id]
                first_excel_row = current_row
                last_excel_row = current_row + len(owner_rows) - 1

                for clean_row in owner_rows:
                    for column, key in enumerate(keys, start=1):
                        value = self._value_for(
                            key=key,
                            row=clean_row,
                            sequence=sequence,
                            excel_row=current_row,
                            property_first_row=first_excel_row,
                            keys=keys,
                        )
                        cell = sheet.cell(row=current_row, column=column, value=value)
                        cell.border = border
                        cell.alignment = self._alignment_for(key)
                        if key in {"total_area", "ownership_area", "share"}:
                            cell.number_format = "0.00"
                        elif key == "living_area":
                            cell.number_format = "0.######"
                    current_row += 1
                    sequence += 1

                if len(owner_rows) > 1:
                    # Merge property-level cells across owners instead of
                    # duplicating values that conceptually belong to the unit.
                    for column, key in enumerate(keys, start=1):
                        if EXPORT_FIELDS_BY_KEY[key].property_level:
                            sheet.merge_cells(
                                start_row=first_excel_row,
                                start_column=column,
                                end_row=last_excel_row,
                                end_column=column,
                            )
                            sheet.cell(first_excel_row, column).alignment = self._alignment_for(key)

                if "total_area" in keys:
                    if total_area_data_start is None:
                        total_area_data_start = first_excel_row
                    total_area_data_end = last_excel_row

        if total_area_data_start is not None and total_area_data_end is not None:
            total_row = current_row + 1
            sheet.cell(total_row, 1, "Усього").font = Font(bold=True)
            total_area_col = keys.index("total_area") + 1
            letter = get_column_letter(total_area_col)
            sheet.cell(
                total_row,
                total_area_col,
                f"=SUM({letter}{total_area_data_start}:{letter}{total_area_data_end})",
            ).font = Font(bold=True)
            sheet.cell(total_row, total_area_col).number_format = "0.00"

        self._set_widths(sheet, keys)
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)

    @staticmethod
    def _sheet_title(registry: RegistryDetails) -> str:
        """Build a compact Excel-safe worksheet title from the query address."""
        title = registry.query_address or "Реєстр"
        for prefix in ("м.Київ, ", "місто Київ, "):
            if title.startswith(prefix):
                title = title[len(prefix) :]
        title = title.replace("вулиця ", "").replace("будинок ", "")
        for char in "[]:*?/\\":
            title = title.replace(char, "-")
        return (title[:31] or "Реєстр").strip()

    @staticmethod
    def _alignment_for(key: str) -> Alignment:
        if key in {"owner_name", "address", "right_record_numbers", "edessb_identifier"}:
            return Alignment(horizontal="left", vertical="center", wrap_text=True)
        return Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _set_widths(sheet, keys: tuple[str, ...]) -> None:
        widths = {
            "unit_number": 16,
            "property_type": 16,
            "owner_no": 8,
            "owner_name": 42,
            "ownership_area": 18,
            "share": 12,
            "total_area": 18,
            "living_area": 18,
            "registry_number": 22,
            "right_record_numbers": 24,
            "right_type": 20,
            "joint_ownership_type": 22,
            "edessb_identifier": 34,
            "address": 48,
            "owner_type": 20,
            "legal_entity_code": 16,
            "registration_country": 20,
        }
        for index, key in enumerate(keys, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(key, 18)

    @staticmethod
    def _value_for(
        *,
        key: str,
        row: RegistryRow,
        sequence: int,
        excel_row: int,
        property_first_row: int,
        keys: tuple[str, ...],
    ) -> object:
        if key == "unit_number":
            return row.unit_number
        if key == "property_type":
            return row.property_type
        if key == "owner_no":
            return sequence
        if key == "owner_name":
            return owner_display_name(row)
        if key == "share":
            if row.share.denominator == 1:
                return row.share.numerator
            return f"={row.share.numerator}/{row.share.denominator}"
        if key == "ownership_area":
            return XlsxExporter._ownership_area_formula(
                row, excel_row, property_first_row, keys
            )
        if key == "total_area":
            return float(row.total_area)
        if key == "living_area":
            return float(row.living_area) if row.living_area is not None else None
        if key == "registry_number":
            return row.registry_number
        if key == "right_record_numbers":
            return ", ".join(row.right_record_numbers)
        if key == "right_type":
            return row.right_type
        if key == "joint_ownership_type":
            return row.joint_ownership_type
        if key == "edessb_identifier":
            return row.edessb_identifier
        if key == "address":
            return row.address
        if key == "owner_type":
            return row.owner_type
        if key == "legal_entity_code":
            return row.legal_entity_code
        if key == "registration_country":
            return row.registration_country
        raise KeyError(key)

    @staticmethod
    def _ownership_area_formula(
        row: RegistryRow,
        excel_row: int,
        property_first_row: int,
        keys: tuple[str, ...],
    ) -> str:
        """Build a formula from visible columns, falling back to literal values."""
        # Referencing selected workbook columns keeps exported calculations live.
        # When a dependency column is omitted, embed the exact domain value so
        # ``ownership_area`` remains correct for any valid field selection.
        if "total_area" in keys:
            total_col = get_column_letter(keys.index("total_area") + 1)
            area_expression = f"{total_col}{property_first_row}"
        else:
            area_expression = str(row.total_area)

        if "share" in keys:
            share_col = get_column_letter(keys.index("share") + 1)
            share_expression = f"{share_col}{excel_row}"
        elif row.share.denominator == 1:
            share_expression = str(row.share.numerator)
        else:
            share_expression = f"({row.share.numerator}/{row.share.denominator})"
        return f"={area_expression}*{share_expression}"
